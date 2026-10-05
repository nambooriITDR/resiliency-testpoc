from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent


def add_list(doc, items):
    for item in items:
        doc.add_paragraph(f'- {item}')


def add_code_block(doc, label, text):
    doc.add_heading(label, level=2)
    for line in text.splitlines():
        doc.add_paragraph(line)


def read_file(path):
    return Path(path).read_text(encoding='utf-8')


doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

p = doc.add_paragraph()
r = p.add_run('Resiliency Test POC')
r.bold = True
r.font.size = Pt(22)

doc.add_paragraph('Disaster Recovery and Cyber Recovery Validation in a Dockerized Environment')
doc.add_paragraph('Prepared for publication and print-ready release.')

doc.add_heading('1. Executive Summary', level=1)
for paragraph in [
    'This project demonstrates a practical proof-of-concept for Disaster Recovery as a Service (DRaaS) and Cyber Recovery as a Service (CRaaS). It simulates a workload, creates a backup, applies a destructive event, and restores the environment from the backup snapshot.',
    'The environment is containerized and uses PostgreSQL, Flask, MinIO, Restic, and OpenSearch. The purpose is to validate operational recovery in a lightweight local lab that can be demonstrated to technical and business audiences.',
    'The key test objective is to prove that both the application data and file set survive a ransomware-style event and can be fully recovered without manual data loss.'
]:
    doc.add_paragraph(paragraph)

doc.add_heading('2. Project Objective', level=1)
add_list(doc, [
    'Generate workload data through the application API.',
    'Create a database dump and a Restic backup.',
    'Simulate a destructive attack by deleting files and dropping schema objects.',
    'Restore the snapshot and replay the SQL dump.',
    'Confirm that both the application and database are recovered.'
])

doc.add_heading('3. Architecture Overview', level=1)
add_list(doc, [
    'PostgreSQL 15 - database layer for application records.',
    'Flask API - workload and data access layer.',
    'MinIO - object storage repository for Restic snapshots.',
    'Restic - backup and restore engine for the database dump and app files.',
    'OpenSearch + Dashboards - monitoring and observability layer.'
])

doc.add_heading('4. Workflow', level=1)
for step in [
    '1. Create records through the application endpoint.',
    '2. Dump the database and store it in /tmp/pgdump.sql.',
    '3. Backup the dump and app files with Restic.',
    '4. Delete the app files and drop the PostgreSQL public schema.',
    '5. Restore the snapshot and replay the SQL dump into PostgreSQL.',
    '6. Validate that the original records reappear in the application and database.'
]:
    doc.add_paragraph(step)

doc.add_heading('5. File Content and Explanations', level=1)

file_sections = {
    'docker-compose.yml': read_file(ROOT / 'docker-compose.yml'),
    'flask/app.py': read_file(ROOT / 'flask' / 'app.py'),
    'flask/Dockerfile': read_file(ROOT / 'flask' / 'Dockerfile'),
    'flask/requirements.txt': read_file(ROOT / 'flask' / 'requirements.txt'),
    'Scripts/backup.sh': read_file(ROOT / 'Scripts' / 'backup.sh'),
    'Scripts/ransom.sh': read_file(ROOT / 'Scripts' / 'ransom.sh'),
    'Scripts/restore.sh': read_file(ROOT / 'Scripts' / 'restore.sh'),
    'generate_final_word.py': read_file(ROOT / 'generate_final_word.py')
}

for name, content in file_sections.items():
    add_code_block(doc, name, content)

    if name == 'docker-compose.yml':
        doc.add_paragraph('This file defines the full stack. PostgreSQL hosts the workload database, a Flask container runs the API, MinIO stores backup content, Restic snapshots the database dump and file set, and OpenSearch provides monitoring and dashboards.')
    elif name == 'flask/app.py':
        doc.add_paragraph('This Python file is the core application logic. It creates the posts table if needed, exposes a create endpoint, exposes a list endpoint, and ensures recovery-safe checks after a schema drop.')
    elif name == 'Scripts/backup.sh':
        doc.add_paragraph('This script performs the backup: it dumps the database and sends both the dump and file set to the Restic repository.')
    elif name == 'Scripts/ransom.sh':
        doc.add_paragraph('This script simulates a destructive cyber event by deleting files and dropping the application schema. It represents a ransomware or corruption scenario.')
    elif name == 'Scripts/restore.sh':
        doc.add_paragraph('This script restores the latest snapshot and replays the SQL dump back into PostgreSQL. The key fix is that the SQL dump is often restored under a nested path such as /tmp/restore/tmp/pgdump.sql.')
    elif name == 'generate_final_word.py':
        doc.add_paragraph('This script creates a printable Word document summary of the project so the outcome can be shared as formal project documentation.')


doc.add_heading('6. Tested Outcome', level=1)
for paragraph in [
    'The full recovery flow was executed and validated in the current environment. The application successfully accepted workload data, the backup was created, the schema was dropped to simulate an attack, and the snapshot was restored successfully.',
    'Verified create request: StatusCode 200; Content = {"content":"demo-record-101","status":"ok"}',
    'Verified backup creation: snapshot 7cebcc71 saved',
    'Verified destructive event: NOTICE: drop cascades to table posts, DROP SCHEMA, CREATE SCHEMA',
    'Verified restore: CREATE TABLE, COPY 2',
    'Verified final recovery: StatusCode 200; Content = [[1,"demo-record-001"],[2,"demo-record-101"]]',
    'Database validation: id | content; 1 | demo-record-001; 2 | demo-record-101; (2 rows)' 
]:
    doc.add_paragraph(paragraph)


doc.add_heading('7. Technical Findings', level=1)
add_list(doc, [
    'The app must create the posts table if the schema was removed after a cyber event.',
    'The transaction needs to be rolled back before re-creating the schema state after a failed operation.',
    'The restore path for the SQL dump is often nested inside /tmp/restore, so exact file detection is required.',
    'Docker Compose shows an obsolete version warning, but the stack still starts correctly under the current Docker Compose version.',
    'This project is suitable for demo, stakeholder education, and technical validation.'
])

doc.add_heading('8. Conclusion', level=1)
for paragraph in [
    'This POC proves that a small but realistic workload can be backed up, destroyed, restored, and validated in a controlled environment.',
    'The results confirm the design objective: the environment survived a simulated ransomware-style event and recovered the database records and application files without permanent loss.',
    'The project is therefore suitable for project documentation, stakeholder presenting, and operational readiness demonstrations.'
]:
    doc.add_paragraph(paragraph)

doc.save(ROOT / 'Resiliency_Test_POC_Document.docx')
print('Created: Resiliency_Test_POC_Document.docx')
